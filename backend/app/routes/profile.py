"""Your profile (from your CV), the standard CV file, and settings."""
import base64
import io

from docx import Document
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.auth import current_profile
from app.database import get_db
from app.defaults import DEFAULT_SETTINGS, user_settings
from app.logic import ai
from app.logic.cv import docx_response, standard_cv_docx
from app.logic.skills import canonical_skill
from app.models import Profile
from app.serialize import to_dict

router = APIRouter(prefix="/api", tags=["profile"])

# Facts not on the CV: email + phone go on the standard CV, visa need is used by the score.
NEEDED_FACTS = {"email": "Email", "phone": "Phone", "visa_required": "Visa sponsorship needed?"}


def profile_dict(p: Profile) -> dict:
    data = to_dict(p)
    data.pop("settings")
    data["missing_facts"] = [label for field, label in NEEDED_FACTS.items() if getattr(p, field) in (None, "")]
    return data


class ProfileUpdate(BaseModel):
    name: str | None = None
    location: str | None = None
    email: str | None = None
    phone: str | None = None
    linkedin_url: str | None = None
    github_url: str | None = None
    headline: str | None = None
    summary: str | None = None
    education: list[str] | None = None
    years_experience: float | None = None
    target_roles: list[str] | None = None
    skills: list[dict] | None = None        # [{"name": "Python", "years": 2.5}]
    languages: list[dict] | None = None
    experience: list[dict] | None = None    # [{"company", "location", "role", "dates", "bullets": [...]}]
    projects: list[dict] | None = None      # [{"name", "tech", "description", "skills", "bullets"}]
    visa_required: bool | None = None
    highlights: list[dict] | None = None    # [{"lead": "...", "text": " — ..."}]
    availability: str | None = None
    career_breaks: list[dict] | None = None  # [{"dates": "December 2024 – August 2025", "text": "..."}]


@router.get("/profile")
def get_profile(profile: Profile = Depends(current_profile)):
    return profile_dict(profile)


@router.patch("/profile")
def update_profile(body: ProfileUpdate, profile: Profile = Depends(current_profile), db: Session = Depends(get_db)):
    changes = body.model_dump(exclude_unset=True)
    if "skills" in changes:
        seen, clean = set(), []
        for s in changes["skills"]:
            name = canonical_skill(str(s.get("name", "")))
            if name and name not in seen:
                seen.add(name)
                clean.append({"name": name, "years": s.get("years")})
        changes["skills"] = clean
    for field, value in changes.items():
        setattr(profile, field, value)
    db.commit()
    return profile_dict(profile)


class CvFile(BaseModel):
    filename: str
    content_base64: str = Field(max_length=14_000_000)  # ≈ 10 MB file


def cv_text(filename: str, data: bytes) -> str:
    """Plain text of a PDF or Word CV (tables included, e.g. the date column)."""
    if filename.lower().endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if filename.lower().endswith(".docx"):
        doc = Document(io.BytesIO(data))
        cells = [p.text for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
        return "\n".join([p.text for p in doc.paragraphs] + cells)
    raise HTTPException(422, "Upload a .pdf or .docx file.")


@router.post("/profile/cv-import")
def import_cv(body: CvFile, profile: Profile = Depends(current_profile)):
    """Read an uploaded CV. Returns its text and, with AI configured, a DRAFT of the profile fields.
    Nothing is saved here: you review the draft on the My CV page and click Save."""
    try:
        text = cv_text(body.filename, base64.b64decode(body.content_base64))
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "Could not read this file. Is it a normal (not scanned or password-protected) PDF or Word file?") from None
    if len(text.strip()) < 200:
        raise HTTPException(422, "Almost no text found in this file (a scanned image?). Paste your CV text instead.")
    if not ai.available():
        return {"text": text, "draft": None,
                "message": "Text extracted. Reading it into the fields needs AI (AI_API_KEY in .env), so copy what you need into the fields below."}
    try:
        draft = ai.extract_cv(text)
    except ai.AIError as exc:
        return {"text": text, "draft": None, "message": str(exc)}
    draft["education"] = [" | ".join(x for x in (e["degree"], e["school"], e["dates"]) if x) for e in draft["education"]]
    draft["skills"] = [{"name": canonical_skill(s) or s, "years": None} for s in draft["skills"]]
    return {"text": text, "draft": draft, "message": "AI read your CV into the fields below. Check every field, then Save."}


@router.get("/cv.docx")
def standard_cv(profile: Profile = Depends(current_profile)):
    """The standard CV: the one file used for every application."""
    return docx_response(standard_cv_docx(profile), profile, None, "CV")


@router.get("/settings")
def get_settings(profile: Profile = Depends(current_profile)):
    return user_settings(profile)


@router.patch("/settings")
def update_settings(body: dict, profile: Profile = Depends(current_profile), db: Session = Depends(get_db)):
    unknown = set(body) - set(DEFAULT_SETTINGS)
    if unknown:
        raise HTTPException(422, f"Unknown settings: {', '.join(sorted(unknown))}")
    if "min_score" in body and not 0 <= float(body["min_score"]) <= 100:
        raise HTTPException(422, "min_score must be between 0 and 100.")
    if "daily_start_time" in body:
        import re

        if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", str(body["daily_start_time"])):
            raise HTTPException(422, "daily_start_time must look like 08:00.")
    if "search_queries" in body:
        keywords = list(dict.fromkeys(" ".join(str(k).split()) for k in body["search_queries"]))  # tidy + no repeats
        body["search_queries"] = [k for k in keywords if k]
        if not body["search_queries"]:
            raise HTTPException(422, "Add at least one job role to search for.")
    if "max_runtime_hours" in body and not 0 < float(body["max_runtime_hours"]) <= 6:
        raise HTTPException(422, "max_runtime_hours must be between 0 and 6.")
    profile.settings = {**(profile.settings or {}), **body}
    db.commit()
    return user_settings(profile)
