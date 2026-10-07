"""Makes a tailored CV (Word .docx) for one job.

Tailoring = ORDER and EMPHASIS only. Everything comes from the profile (your real CV):
- "Key skills for this role": your skills that the job asks for, listed first
- experience bullets and projects sorted so the most relevant ones come first
Nothing is added that is not already on your CV.
"""
import io

from docx import Document
from docx.shared import Pt
from fastapi.responses import Response

from app.logic.skills import find_skills


def _relevance(text: str, wanted: set[str]) -> int:
    return len(set(find_skills(text)) & wanted)


def tailored_cv_docx(profile, job, job_skills: list[str], version: dict | None = None) -> bytes:
    """version = one of defaults.CV_VERSIONS: its headline and skill focus order (only skills you really have)."""
    wanted = set(job_skills)
    focus = (version or {}).get("focus", [])
    my_skills = sorted((s["name"] for s in profile.skills), key=lambda s: focus.index(s) if s in focus else len(focus))
    matching = [s for s in my_skills if s in wanted]
    others = [s for s in my_skills if s not in wanted]

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(profile.name, level=0)
    headline = (version or {}).get("headline") or profile.headline
    if headline:
        doc.add_paragraph(headline)
    contact = " | ".join(x for x in (profile.location, profile.email, profile.phone, profile.linkedin_url, profile.github_url) if x)
    doc.add_paragraph(contact)

    if profile.summary:
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(profile.summary)

    doc.add_heading("Key Skills", level=1)
    if matching:
        doc.add_paragraph("Most relevant for this role: " + ", ".join(matching))
    doc.add_paragraph(("Also: " if matching else "") + ", ".join(others))

    doc.add_heading("Professional Experience", level=1)
    for job_entry in profile.experience or []:
        doc.add_paragraph(f"{job_entry['role']} — {job_entry['company']}, {job_entry['location']} ({job_entry['dates']})").runs[0].bold = True
        for bullet in sorted(job_entry.get("bullets", []), key=lambda b: -_relevance(b, wanted)):
            doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("Key Projects", level=1)
    projects = sorted(profile.projects, key=lambda p: -len(set(p.get("skills", [])) & wanted))
    for project in projects[:4]:
        title = project["name"] + (f"  ({project['tech']})" if project.get("tech") else "")
        doc.add_paragraph(title).runs[0].bold = True
        for bullet in sorted(project.get("bullets", []), key=lambda b: -_relevance(b, wanted)):
            doc.add_paragraph(bullet, style="List Bullet")

    if profile.education:
        doc.add_heading("Education", level=1)
        for line in profile.education:
            doc.add_paragraph(line)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def text_docx(text: str) -> bytes:
    """A simple .docx for the cover letter."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    for block in text.split("\n"):
        doc.add_paragraph(block)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def docx_response(data: bytes, profile, job, kind: str) -> Response:
    """Send a .docx as a download named like 'Arpita_Chaudhari_CV_Celonis.docx'."""
    company = "".join(c if c.isalnum() else "_" for c in job.company)[:40]
    name = f"{profile.name.replace(' ', '_')}_{kind}_{company}.docx"
    return Response(data, media_type=DOCX, headers={"Content-Disposition": f'attachment; filename="{name}"'})
