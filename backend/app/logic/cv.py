"""The standard CV (Word .docx) used for EVERY application, and the cover letter .docx.

A4, header with headline + one-line contact + availability, then profile, key highlights, skills, experience
(MM/YYYY date column), projects, education (redesign approved by the user 2026-10-23; no visa line, no photo / date of birth).
Everything comes from the profile (your real CV). Nothing is added, nothing is tailored per job.
"""
import io
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from fastapi.responses import Response

ACCENT, GREY = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x5F, 0x63, 0x68)
LEFT = Cm(3.2)  # width of the left (dates / labels) column

# Skill groups for the "Technical skills" section; a skill not listed here goes to "Other".
SKILL_GROUPS = {
    "Programming": ["Python", "TypeScript", "JavaScript", "FastAPI", "Django", "Node.js", "Express.js", "Angular", "React"],
    "Industrial & OT": ["SCADA", "WebSCADA", "OPC", "OPC UA", "Modbus", "FINS", "PLC", "IIoT", "21 CFR Part 11", "LabVIEW"],
    "Databases & cloud": ["MySQL", "PostgreSQL", "MongoDB", "AWS"],
    "Tools & practices": ["REST APIs", "JWT", "RBAC", "PyInstaller", "Windows Services", "Docker", "Git", "Postman", "Agile"],
}
# How a stored (canonical) skill is written on the CV. She connects to LabVIEW but doesn't program in it (user, 2026-10-23).
SKILL_DISPLAY = {"LabVIEW": "LabVIEW connectivity", "PLC": "PLC data acquisition"}  # she reads PLC data; no PLC programming
MONTHS = {m: i + 1 for i, m in enumerate(["january", "february", "march", "april", "may", "june", "july",
                                         "august", "september", "october", "november", "december"])}


def eu_dates(text: str) -> str:
    """'April 2026 – Present' → '04/2026 – present' (the usual format on German CVs)."""
    text = re.sub(r"\b([A-Za-z]+) (\d{4})\b", lambda m: f"{MONTHS[m[1].lower()]:02d}/{m[2]}" if m[1].lower() in MONTHS else m[0], text)
    return re.sub(r"\bPresent\b", "present", text)


def start_key(dates: str) -> tuple[int, int]:
    """'December 2024 – August 2025' / '12/2024 – 08/2025' → (2024, 12): to sort the timeline."""
    m = re.search(r"(\d{1,2})/(\d{4})", eu_dates(dates))
    return (int(m[2]), int(m[1])) if m else (0, 0)


def _run(par, text: str, size: float = 10, bold=False, italic=False, color=None):
    run = par.add_run(text)
    run.font.size, run.bold, run.italic = Pt(size), bold, italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _heading(doc, text: str):
    par = doc.add_paragraph()
    par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(7), Pt(2.5)
    par.paragraph_format.keep_with_next = True  # never a heading alone at the bottom of a page
    _run(par, text.upper(), 10.5, bold=True, color=ACCENT)
    border = OxmlElement("w:pBdr")  # thin line under the heading
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F3A5F")):
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    par._p.get_or_add_pPr().append(border)


def _row(doc, left: str):
    """A line with grey text in the left column (tab stop) and a hanging indent: wrapped text stays in the right column.
    Plain paragraphs, not a table: Word and Pages both break them across pages normally."""
    par = doc.add_paragraph()
    fmt = par.paragraph_format
    fmt.left_indent, fmt.first_line_indent, fmt.space_after = LEFT, -LEFT, Pt(1)
    fmt.tab_stops.add_tab_stop(LEFT)
    _run(par, left, 9.3, color=GREY)
    par.add_run("\t")
    return par


def _indented(doc):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = LEFT
    return par


def _bullet(doc, text: str, indent=Cm(0)):
    par = doc.add_paragraph(style="List Bullet")
    fmt = par.paragraph_format
    fmt.left_indent, fmt.first_line_indent, fmt.space_after = indent + Cm(0.5), -Cm(0.5), Pt(1.5)
    if text:
        _run(par, text, 9.2)
    return par


def education_parts(line: str) -> tuple[str, str, str]:
    """'MSc Computer Science | Greenwich University, London | January 2022 – January 2023' → (degree, school, dates).
    Parts may be missing; the one with a year is the dates."""
    parts = [p.strip() for p in line.split("|") if p.strip()]
    dates = next((p for p in parts if re.search(r"\d{4}", p)), "")
    rest = [p for p in parts if p != dates] + ["", ""]
    return rest[0], rest[1], dates


def degree_text(line: str) -> str:
    """One education line as plain text: 'MSc Computer Science, Greenwich University, London'."""
    degree, school, _ = education_parts(line)
    return f"{degree}, {school}" if school else degree


def phone_text(phone: str | None) -> str | None:
    """'+919812345678' → '+91 98123 45678' (easier to read; other formats stay as typed)."""
    digits = re.sub(r"\D", "", phone or "")
    return f"+91 {digits[2:7]} {digits[7:]}" if (phone or "").startswith("+91") and len(digits) == 12 else phone


def skill_groups(skills: list[str]) -> list[tuple[str, list[str]]]:
    groups = [(name, [s for s in members if s in skills]) for name, members in SKILL_GROUPS.items()]
    known = {s for members in SKILL_GROUPS.values() for s in members}
    groups.append(("Other", [s for s in skills if s not in known]))
    return [(name, [SKILL_DISPLAY.get(s, s) for s in members]) for name, members in groups if members]


TEAL, DARK = RGBColor(0x0F, 0x76, 0x6E), RGBColor(0x22, 0x22, 0x22)
# Technology words set in bold inside bullets, so a 6-second skim still catches the stack (the text itself is unchanged).
KEYWORDS = [
    "Python", "FastAPI", "Django", "Angular", "TypeScript", "Node.js/Express.js", "Node.js", "Express.js", "MongoDB",
    "MySQL", "PostgreSQL", "React", "AWS EC2", "JWT", "RBAC", "OPC-UA", "OPC UA", "Modbus TCP/RTU", "FINS Serial",
    "FINS Ethernet", "21 CFR Part 11", "WebSCADA", "SCADA", "PyInstaller", "NSSM", "Windows services", "LabVIEW",
    "Playwright", "Kepware", "RESTful APIs", "Postman", "~30%", "zero licensing cost", "Agile",
]
KEYWORD_RE = re.compile("(" + "|".join(re.escape(k) for k in sorted(KEYWORDS, key=len, reverse=True)) + ")")


def _rich(par, text: str, size: float = 9.4):
    """Text with the technology words in bold."""
    for i, part in enumerate(KEYWORD_RE.split(text)):
        if part:
            _run(par, part, size, bold=(i % 2 == 1), color=DARK)


CV_FONT = "Arial Narrow"


def _link(par, text: str, url: str, size: float, color):
    """A clickable link (kept in the PDF Pages exports), styled like the text around it."""
    rel = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, attrs in (("w:color", {"w:val": str(color)}), ("w:sz", {"w:val": str(int(size * 2))})):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    par._p.append(link)


def _rule(par, size: str = "6"):
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", size), ("space", "1"), ("color", "1F3A5F")):
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    par._p.get_or_add_pPr().append(border)


def standard_cv_docx(profile) -> bytes:
    """Header (name, headline, contact, availability) → profile → key highlights → skills → experience → projects → education."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.6)
    section.top_margin = section.bottom_margin = Cm(0.95)
    normal = doc.styles["Normal"]
    # Not Calibri: Pages exports its "ft"/"tt" ligatures as single glyphs, so application systems read "So=ware Engineer"
    # (checked 2026-10-23). Arial Narrow extracts cleanly and keeps the CV at 2 pages.
    normal.font.name, normal.font.size = CV_FONT, Pt(10)
    normal.paragraph_format.space_after = Pt(0)

    _run(doc.add_paragraph(), profile.name, 25, bold=True, color=ACCENT)
    if profile.headline:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(3)
        _run(par, profile.headline, 12.5, bold=True, color=TEAL)
    for line in ([(profile.location, None), (phone_text(profile.phone), None),
                  (profile.email, f"mailto:{profile.email}" if profile.email else None)],
                 [(u.removeprefix("https://"), u if u.startswith("http") else f"https://{u}")
                  for u in (profile.linkedin_url, profile.github_url) if u]):
        items = [(t, u) for t, u in line if t]
        if items:
            par = doc.add_paragraph()
            for i, (label, url) in enumerate(items):
                if i:
                    _run(par, "  ·  ", 9.5, color=GREY)
                if url:
                    _link(par, label, url, 9.5, GREY)
                else:
                    _run(par, label, 9.5, color=GREY)
    languages = ", ".join(f"{x['language']} ({x['level']})" if x.get("level") else x["language"] for x in profile.languages or [])
    extra = "  ·  ".join(x for x in (getattr(profile, "availability", None), languages) if x)
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(2)
    if extra:
        _run(par, extra, 9.5, bold=True, color=ACCENT)
    _rule(par, "12")

    if profile.summary:
        _heading(doc, "Profile")
        par = doc.add_paragraph()
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _rich(par, profile.summary, 10)

    highlights = getattr(profile, "highlights", None) or []
    if highlights:
        _heading(doc, "Key highlights")
        for h in highlights:
            par = _bullet(doc, "")
            _run(par, h.get("lead", ""), 9.6, bold=True, color=ACCENT)
            _run(par, h.get("text", ""), 9.6, color=DARK)

    groups = skill_groups([s["name"] for s in profile.skills or []])
    if groups:  # early: recruiters and ATS look for the stack first
        _heading(doc, "Technical skills")
        for name, members in groups:
            _run(_row(doc, name), ", ".join(members), 9.8, color=DARK)

    _heading(doc, "Professional experience")
    timeline = [dict(e) for e in profile.experience or []] + \
        [{**b, "break": True} for b in getattr(profile, "career_breaks", None) or []]
    timeline.sort(key=lambda e: start_key(e.get("dates", "")), reverse=True)  # newest first, breaks in between
    for entry in timeline:
        if entry.get("break"):
            par = _row(doc, eu_dates(entry.get("dates", "")))
            par.paragraph_format.space_before = Pt(4)
            _run(par, entry.get("text", ""), 10, italic=True, color=GREY)
            continue
        par = _row(doc, eu_dates(entry.get("dates", "")))
        par.paragraph_format.space_before, par.paragraph_format.keep_with_next = Pt(4), True
        _run(par, entry["role"], 10.5, bold=True, color=DARK)
        par = _indented(doc)
        par.paragraph_format.keep_with_next = True
        _run(par, " · ".join(x for x in (entry.get("company"), entry.get("location")) if x), 9.8, italic=True, color=TEAL)
        for bullet in entry.get("bullets", []):
            _rich(_bullet(doc, "", LEFT), bullet)

    if profile.projects:
        _heading(doc, "Key projects")
        for project in profile.projects:
            par = doc.add_paragraph()
            par.paragraph_format.space_before, par.paragraph_format.keep_with_next = Pt(5), True
            _run(par, project["name"], 10.5, bold=True, color=DARK)
            if project.get("tech"):
                par = doc.add_paragraph()
                par.paragraph_format.keep_with_next = True
                _run(par, project["tech"], 8.8, italic=True, color=TEAL)
            for bullet in project.get("bullets", []):
                _rich(_bullet(doc, ""), bullet, 9.2)

    if profile.education:
        _heading(doc, "Education")
        for line in profile.education:  # "degree | school | dates", newest first
            degree, school, dates = education_parts(line)
            par = _row(doc, eu_dates(dates))
            _run(par, degree, 10.3, bold=True, color=DARK)
            if school:
                _run(par, f" — {school}", 9.8, color=GREY)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def text_docx(text: str) -> bytes:
    """A simple .docx for the cover letter."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    for block in text.split("\n"):
        doc.add_paragraph(block)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def docx_response(data: bytes, profile, job, kind: str) -> Response:
    """Send a .docx as a download: 'Arpita_Chaudhari_CV.docx', or 'Arpita_Chaudhari_Cover_Letter_Celonis.docx' with a job."""
    company = "_" + "".join(c if c.isalnum() else "_" for c in job.company)[:40] if job is not None else ""
    name = f"{profile.name.replace(' ', '_')}_{kind}{company}.docx"
    return Response(data, media_type=DOCX, headers={"Content-Disposition": f'attachment; filename="{name}"'})
