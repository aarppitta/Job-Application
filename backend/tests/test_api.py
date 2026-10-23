"""End-to-end tests through the API (uses the test database)."""
import io
import time

import pytest

from docx import Document

from app.defaults import DEFAULT_PROFILE
from tests.test_logic import SCADA_JD

JOB = {"title": "SCADA Software Developer", "company": "Muster Automation GmbH", "description": SCADA_JD,
       "location": "Stuttgart, Germany", "application_url": "https://job-boards.greenhouse.io/muster/jobs/1"}


def docx_text(content: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(content)).paragraphs if p.text.strip()]


def test_login_only_master_user(client):
    assert client.post("/api/auth/login", json={"username": "arpita", "password": "wrong"}).status_code == 401
    assert client.post("/api/auth/register", json={"username": "x", "password": "y"}).status_code in (404, 405)
    profile = client.get("/api/profile").json()
    assert profile["name"] == "Arpita Chaudhari"
    assert profile["experience"][0]["company"] == "Kevision Systems"
    assert profile["missing_facts"] == ["Email", "Phone", "Visa sponsorship needed?"]
    assert client.patch("/api/settings", json={"min_score": 150}).status_code == 422


def test_import_dedup_and_jd_versions(client):
    first = client.post("/api/jobs/import", json={**JOB, "source": "linkedin"}).json()
    job = client.get(f"/api/jobs/{first['id']}").json()
    assert job["score"]["recommendation"] == "APPLY" and job["status"] == "QUALIFIED"

    same = client.post("/api/jobs/import", json={**JOB, "title": "Software Developer SCADA", "source": "company"}).json()
    assert same["outcome"] == "duplicate" and same["id"] == first["id"]

    changed = client.post("/api/jobs/import", json={**JOB, "description": SCADA_JD.replace("3+", "4+"), "source": "linkedin"}).json()
    assert changed["outcome"] == "updated"
    assert [v["version"] for v in client.get(f"/api/jobs/{first['id']}/versions").json()] == [1, 2]
    assert client.get("/api/jobs", params={"keyword": "scada", "visa": True}).json()["total"] == 1


def test_standard_cv_and_cover_letter_buttons(client):
    job_id = client.post("/api/jobs/import", json=JOB).json()["id"]
    other_id = client.post("/api/jobs/import", json={**JOB, "title": "Python Developer", "company": "Other GmbH",
                                                      "description": "Python FastAPI developer. " + SCADA_JD[::-1],
                                                      "application_url": "https://example.com/jobs/other"}).json()["id"]
    client.patch("/api/profile", json={"email": "me@example.com", "visa_required": True})

    cv = client.get(f"/api/jobs/{job_id}/cv.docx")
    assert cv.status_code == 200 and 'filename="Arpita_Chaudhari_CV.docx"' in cv.headers["content-disposition"]
    # One standard CV: the same file for every job and for /api/cv.docx, no company name, nothing tailored
    lines = docx_text(cv.content)
    assert docx_text(client.get(f"/api/jobs/{other_id}/cv.docx").content) == lines == docx_text(client.get("/api/cv.docx").content)
    assert "Email\tme@example.com" in lines and not any(line.startswith("Work permit") for line in lines)   # visa need: not on the CV
    assert "04/2026 – present\tSystem Software Engineer — Industrial Automation" in lines          # German date column
    assert not any("Muster" in line for line in lines)
    real = {b for e in DEFAULT_PROFILE["experience"] for b in e["bullets"]} | {b for p in DEFAULT_PROFILE["projects"] for b in p["bullets"]}
    assert real <= set(lines)                       # every real bullet, unchanged
    assert "01/2022 – 01/2023\tMSc Computer Science — Greenwich University, London, United Kingdom" in lines
    section = Document(io.BytesIO(cv.content)).sections[0]
    assert round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7   # A4

    edited = [{**DEFAULT_PROFILE["projects"][0], "bullets": ["A bullet I edited myself."]}]
    assert client.patch("/api/profile", json={"projects": edited}).json()["projects"] == edited       # CV edits via the API
    assert "A bullet I edited myself." in docx_text(client.get("/api/cv.docx").content)

    letter = client.post(f"/api/jobs/{job_id}/cover-letter").json()
    assert "Muster Automation GmbH" in letter["text"] and "Kevision Systems" in letter["text"] and not letter["warnings"]
    word = client.post(f"/api/jobs/{job_id}/cover-letter.docx", json={"text": letter["text"] + "\nEdited by me."})
    assert "Edited by me." in docx_text(word.content)


def test_mark_applied_keeps_records(client):
    job_id = client.post("/api/jobs/import", json=JOB).json()["id"]
    app = client.post(f"/api/jobs/{job_id}/applied", json={"cover_letter": "Dear team, my edited letter."}).json()
    assert app["status"] == "APPLIED"
    assert client.post(f"/api/jobs/{job_id}/applied", json={}).status_code == 409   # never twice

    [saved] = client.get("/api/applications").json()
    assert saved["cover_letter"] == "Dear team, my edited letter." and saved["jd_version"] == 1 and saved["score"] >= 80
    assert client.get(f"/api/applications/{saved['id']}/cv.docx").status_code == 200
    assert client.get(f"/api/applications/{saved['id']}/cover-letter.docx").status_code == 200

    moved = client.patch(f"/api/applications/{saved['id']}", json={"status": "HR_SCREEN"}).json()
    assert [h["status"] for h in moved["history"]] == ["APPLIED", "HR_SCREEN"]
    assert client.patch(f"/api/applications/{saved['id']}", json={"status": "BOGUS"}).status_code == 422

    counts = client.get("/api/dashboard").json()["counts"]
    assert counts["Applied"] == 1 and counts["Interview"] == 1 and counts["Jobs found (total)"] == 1


def wait_for_end(client, seconds=20):
    for _ in range(seconds * 10):
        run = client.get("/api/automation/status").json()
        if run["status"] not in ("RUNNING", "PAUSED"):
            return run
        time.sleep(0.1)
    raise AssertionError("automation did not finish")


def test_automation_finds_and_scores(client):
    client.patch("/api/settings", json={"sources": ["demo"], "search_queries": ["*"], "repeat_every_minutes": 0})
    assert client.post("/api/automation/start", json={}).status_code == 200
    assert client.post("/api/automation/start", json={}).status_code == 409   # only one at a time
    run = wait_for_end(client)
    assert run["status"] == "FINISHED" and run["discovered"] == 6 and 1 <= run["qualified"] < 6
    assert any(line["message"].startswith("MATCH:") for line in client.get("/api/automation/logs").json())
    top = client.get("/api/jobs", params={"min_score": 85}).json()["items"]
    assert top and all(j["ats_score"] >= 85 for j in top)

    client.post("/api/automation/start", json={})
    assert wait_for_end(client)["duplicates"] == 6   # second run: nothing new


def test_automation_pause_resume_stop(client):
    client.patch("/api/settings", json={"sources": ["demo"], "search_queries": ["*"] * 50, "repeat_every_minutes": 0})
    client.post("/api/automation/start", json={})
    assert client.post("/api/automation/pause").json()["status"] == "PAUSED"
    assert client.post("/api/automation/resume").json()["status"] == "RUNNING"
    client.post("/api/automation/stop")
    wait_for_end(client)
    time.sleep(1.5)
    run = client.get("/api/automation/status").json()
    assert run["status"] == "STOPPED" and "Stopped by you" in (run["summary"] or "")


def test_today_week_older_tabs_and_labels(client):
    from datetime import UTC, datetime, timedelta

    client.post("/api/jobs/import", json=JOB)  # posted today (no date given = found today)
    old = (datetime.now(UTC) - timedelta(days=10)).isoformat()
    client.post("/api/jobs/import", json={**JOB, "title": "Python Backend Engineer", "company": "Old GmbH",
                                         "description": SCADA_JD + " Python backend.", "posted_at": old,
                                         "application_url": "https://jobs.lever.co/old/1"})
    assert client.get("/api/jobs", params={"view": "today"}).json()["total"] == 1
    assert client.get("/api/jobs", params={"view": "week"}).json()["total"] == 1   # default view
    older = client.get("/api/jobs", params={"view": "older"}).json()["items"]
    assert [j["company"] for j in older] == ["Old GmbH"]
    today = client.get("/api/jobs", params={"view": "today"}).json()["items"][0]
    assert today["match"]["level"] == "good" and today["match"]["text"].startswith("Good match")


def test_interview_prep(client):
    job_id = client.post("/api/jobs/import", json=JOB).json()["id"]
    prep = client.get(f"/api/jobs/{job_id}/interview-prep").json()
    skills = [t["skill"] for t in prep["technical"]]
    assert "OPC UA" in skills and "Modbus" in skills
    assert any(s["skill"] == "OPC UA" and "OPC" in s["where"] for s in prep["stories"])   # a real CV project
    assert any(g["skill"] == "Kubernetes" for g in prep["gaps"])
    assert prep["general_questions"]
    dashboard = client.get("/api/dashboard").json()
    assert any(s["skill"] == "Kubernetes" for s in dashboard["skills_to_learn"])


def test_recruiters_page_and_pitch(client):
    job_id = client.post("/api/jobs/import", json=JOB).json()["id"]   # Stuttgart
    data = client.get("/api/recruiters").json()
    [company] = data["companies"]
    assert company["company"] == "Muster Automation GmbH" and company["jobs"][0]["id"] == job_id
    assert company["links"]["LinkedIn: recruiters"].startswith("https://www.linkedin.com/search/results/people/?keywords=recruiter")
    assert client.get("/api/recruiters", params={"city": "Stuttgart"}).json()["companies"]
    assert client.get("/api/recruiters", params={"city": "Berlin"}).json()["companies"] == []

    p = client.post(f"/api/jobs/{job_id}/pitch", json={"contact_name": "Anna Schmidt"}).json()
    assert p["note"].startswith("Hi Anna,") and len(p["note"]) <= 300 and "relocating to Germany" in p["note"]
    assert "I'm interested in the SCADA Software Developer role at Muster Automation GmbH" in p["note"]
    assert "Kevision Systems" in p["message"] and p["email_subject"].startswith("SCADA Software Developer")
    client.post(f"/api/jobs/{job_id}/applied", json={})
    assert "I applied for" in client.post(f"/api/jobs/{job_id}/pitch", json={}).json()["note"]


def test_job_boards_page(client):
    data = client.get("/api/job-boards", params={"q": "Python Developer"}).json()
    groups = {g["group"]: g["boards"] for g in data["groups"]}
    assert set(groups) == {"Germany", "Cities (startup boards)", "Europe", "Remote in Europe", "Direct employers"}
    remote = {b["name"]: b["kind"] for b in groups["Remote in Europe"]}
    assert remote["Himalayas"] == "automatic" and remote["Wellfound (remote)"] == "link" and len(remote) == 9
    cities = {b["name"]: b for b in groups["Cities (startup boards)"]}
    assert cities["Berlin Startup Jobs"]["kind"] == "automatic" and cities["englishjobs.de"]["kind"] == "link"
    names = {b["name"]: b for b in groups["Germany"]}
    assert names["Arbeitnow"]["kind"] == "automatic"
    assert names["LinkedIn Jobs"]["url"] == "https://www.linkedin.com/jobs/search/?keywords=Python%20Developer&location=Germany"
    assert names["Indeed Germany"]["url"].startswith("https://de.indeed.com/jobs?q=Python+Developer")
    assert any(b["name"] == "Siemens" for b in groups["Direct employers"])


def test_arbeitnow_parsing(monkeypatch):
    from app import sources

    page = {"data": [
        {"title": "Python Backend Engineer", "company_name": "ACME", "description": "<p>Python, FastAPI</p>", "remote": False,
         "url": "https://example.com/job/1", "location": "Berlin", "created_at": 1790000000},
        {"title": "Python Backend Engineer", "company_name": "Far Away", "description": "Python", "remote": False,
         "url": "https://example.com/job/2", "location": "Paris", "created_at": 1790000000},
    ], "links": {"next": None}}

    class Fake:
        def json(self):
            return page

    monkeypatch.setattr(sources, "get_cached", lambda *a, **k: Fake())
    jobs = sources.arbeitnow("Python Engineer", [])
    assert [j["company"] for j in jobs] == ["ACME"]           # only Germany
    assert jobs[0]["posted_at"].year == 2026 and jobs[0]["source"] == "arbeitnow"


def test_job_keywords_setting(client):
    assert client.patch("/api/settings", json={"search_queries": ["", "  "]}).status_code == 422
    saved = client.patch("/api/settings", json={"search_queries": [" Data  Analyst", "Data Analyst", "BI Developer"]}).json()
    assert saved["search_queries"] == ["Data Analyst", "BI Developer"]
    assert client.get("/api/job-boards").json()["keyword"] == "Data Analyst"


class FakePage:
    def __init__(self, text):
        self.text, self.content = text, text.encode()


def test_city_boards_parsing(monkeypatch):
    from app import sources

    rss = """<rss><channel>
      <item><title>Python Backend Engineer // ACME</title><link>https://bsj/1</link>
        <pubDate>Mon, 28 Sep 2026 11:50:22 +0000</pubDate><description>teaser</description></item>
      <item><title>Sales Engineer // Other</title><link>https://bsj/2</link><description>software</description></item>
    </channel></rss>"""
    job_page = """<script type="application/ld+json">{"@type": "JobPosting", "datePosted": "2026-09-27",
      "description": "Python backend work with FastAPI", "hiringOrganization": {"name": "ACME GmbH"}}</script>"""
    monkeypatch.setattr(sources, "get_cached", lambda source, url, params=None: FakePage(rss if "feed" in url else job_page))
    jobs = sources.berlinstartupjobs("Python Backend Engineer", [])
    assert [(j["title"], j["company"]) for j in jobs] == [("Python Backend Engineer", "ACME GmbH")]
    assert jobs[0]["posted_at"].isoformat() == "2026-09-27T00:00:00+00:00" and jobs[0]["location"] == "Berlin, Germany"

    listing = """<a class="card" href="/jobs/anzeige/7"><img alt="x"/><p class="text-card-meta truncate">Robo GmbH</p></div>
      <h3 class="text-card-title">Python Developer</h3></a>"""
    detail = '<h1>Python Developer</h1><span>vor 2 Wochen</span><div class="pinboard-prose max-w-none"><p>Python, Docker</p></div></article>'
    monkeypatch.setattr(sources, "get_cached", lambda source, url, params=None: FakePage(detail if "anzeige" in url else listing))
    jobs = sources.munichstartup("Python Developer", [])
    assert [(j["company"], j["source_url"]) for j in jobs] == [("Robo GmbH", "https://www.munich-startup.de/jobs/anzeige/7")]
    assert jobs[0]["posted_at"] is not None and "Docker" in jobs[0]["description"]


def test_ashby_parsing(monkeypatch):
    from app import sources

    board = {"jobs": [
        {"title": "Software Engineer (f/m/d)", "location": "Berlin", "isListed": True, "isRemote": False,
         "descriptionPlain": "Python, TypeScript", "descriptionHtml": "<p>Python, TypeScript</p>",
         "jobUrl": "https://jobs.ashbyhq.com/acme/1", "applyUrl": "https://jobs.ashbyhq.com/acme/1/application",
         "publishedAt": "2026-10-01T10:00:00.000+00:00"},
        {"title": "Software Engineer", "location": "Paris", "isListed": True, "descriptionPlain": "Python",
         "jobUrl": "https://jobs.ashbyhq.com/acme/2"},                              # not in Germany
        {"title": "Account Executive", "location": "Munich", "isListed": True, "descriptionPlain": "sales",
         "jobUrl": "https://jobs.ashbyhq.com/acme/3"},                              # not the role
    ]}

    class Fake:
        def json(self):
            return board

    monkeypatch.setattr(sources, "get_cached", lambda *a, **k: Fake())
    jobs = sources.ashby("Mid Software Engineer", ["black-forest-labs"])
    assert [(j["company"], j["location"]) for j in jobs] == [("Black Forest Labs", "Berlin")]
    assert jobs[0]["application_url"].endswith("/application") and jobs[0]["posted_at"].day == 1


def test_job_analysis_is_honest(client):
    job_id = client.post("/api/jobs/import", json=JOB).json()["id"]
    a = client.get(f"/api/jobs/{job_id}/analysis").json()
    info = a["job_information"]
    assert info["Company"] == "Muster Automation GmbH" and info["Salary"] == "Not specified"   # not in the JD -> not guessed
    must = {r["item"]: r["status"] for r in a["requirements"]["must"]}
    assert must["Python"] == "strong" and must["OPC UA"] == "strong"
    assert a["scores"]["range"][0] <= a["scores"]["overall"] <= a["scores"]["range"][1]
    assert a["dsa"]["level"] == "Low–Medium" and a["dsa"]["basis"].startswith("Estimated")          # industrial company
    assert a["recommendation"]["decision"] == "APPLY" and a["cv_plan"]["cv_version"] == "STANDARD"
    assert any("L2 LineSync" in q["question"] for q in a["predicted_questions"])                       # questions from YOUR CV
    assert all(k["keyword"] not in {s["name"] for s in DEFAULT_PROFILE["skills"]} for k in a["ats"]["must_not_add"])
    assert a["research"] is None and a["readiness"]["note"].startswith("Estimated")


def test_contacts_health_and_follow_up(client):
    job_id = client.post("/api/jobs/import", json=JOB).json()["id"]
    assert client.get(f"/api/jobs/{job_id}/analysis").json()["health"]["score"] == 10                  # only "JD match"
    client.post(f"/api/jobs/{job_id}/applied", json={})
    rec = client.post(f"/api/jobs/{job_id}/contacts", json={
        "name": "Sarah Müller", "role": "Technical Recruiter", "kind": "recruiter",
        "linkedin_url": "https://www.linkedin.com/in/example"}).json()
    assert rec["tier"] == 1
    assert client.post(f"/api/jobs/{job_id}/contacts", json={"name": "X Y", "role": "r", "kind": "recruiter",
                                                             "linkedin_url": "https://evil.example.com"}).status_code == 422
    msg = client.get(f"/api/contacts/{rec['id']}/message").json()
    assert msg["note"].startswith("Hi Sarah,") and len(msg["note"]) <= 300 and "Muster Automation GmbH" in msg["note"]
    assert msg["email_allowed"] is False                                                                # no public email given
    client.patch(f"/api/contacts/{rec['id']}", json={"contacted": True})
    [app] = client.get("/api/applications").json()
    assert app["status"] == "RECRUITER_CONTACTED" and app["cv_version"] == "STANDARD" and app["cv_version_name"] == "Standard CV" and app["next_follow_up"]
    assert app["health"]["score"] >= 70 and app["snapshot"]["recommendation"]["decision"] == "APPLY"
    for i in range(2):
        client.post(f"/api/jobs/{job_id}/contacts", json={"name": f"Person {i}", "role": "Engineering Manager", "kind": "hiring_manager"})
    too_many = client.post(f"/api/jobs/{job_id}/contacts", json={"name": "One More", "role": "Head", "kind": "department_head"})
    assert too_many.status_code == 409                                                                  # quality over quantity


def test_mock_interview_without_ai_and_strategy(client):
    job_id = client.post("/api/jobs/import", json=JOB).json()["id"]
    pkg = client.get(f"/api/jobs/{job_id}/interview-package").json()
    assert pkg["questions"] and pkg["readiness"]["weakest"] and pkg["ai_available"] is False
    s = client.post(f"/api/jobs/{job_id}/mock", json={}).json()
    first = s["next"]["question"]
    s = client.post(f"/api/mock/{s['id']}/answer", json={"answer": "I built L2 LineSync with FastAPI.", "self_score": 70}).json()
    assert s["answered"] == 1 and s["answers"][0]["question"] == first and s["answers"][0]["grade"]["by"] == "self"
    assert s["answers"][0]["grade"]["ideal_structure"] and s["average"] == 70
    assert client.post(f"/api/jobs/{job_id}/research").status_code == 400                               # no AI key: clear error

    client.post(f"/api/jobs/{job_id}/applied", json={})
    [app] = client.get("/api/applications").json()
    client.patch(f"/api/applications/{app['id']}", json={"status": "REJECTED", "rejection_stage": "TECHNICAL"})
    st = client.get("/api/strategy").json()
    assert st["goal"]["goal_date"] == "2027-01-31" and st["funnel"]["applications"] == 1
    assert st["funnel"]["technical_rounds"] == 1 and st["rejections"]["by_stage"] == {"TECHNICAL": 1}
    assert client.patch(f"/api/applications/{app['id']}", json={"rejection_stage": "NOPE"}).status_code == 422


def test_research_sources_are_verified():
    from app.logic.ai import verify_sources

    data = {"facts": [{"topic": "size", "value": "500", "status": "confirmed", "source_url": "https://real.example/about/"},
                      {"topic": "salary_range", "value": "€90k", "status": "reported", "source_url": "https://made-up.example"}],
            "interview_rounds": [{"order": 1, "name": "HR", "status": "reported", "source_url": "https://made-up.example"}],
            "reported_questions": [{"question": "Q?", "source_url": "https://made-up.example"}],
            "dsa": {"level": "High", "evidence": "x", "source_url": "https://made-up.example"}}
    out = verify_sources(data, ["https://real.example/about"])
    assert out["facts"][0]["status"] == "confirmed"
    assert out["facts"][1]["status"] == "could_not_verify" and out["facts"][1]["source_url"] == ""
    assert out["interview_rounds"][0]["status"] == "estimated" and out["reported_questions"] == []
    assert out["dsa"]["level"] == "Not found"


def test_answer_guides_and_ai_draft(client, monkeypatch):
    job_id = client.post("/api/jobs/import", json=JOB).json()["id"]
    qs = client.get(f"/api/jobs/{job_id}/analysis").json()["predicted_questions"]
    intro = next(q for q in qs if q["question"] == "Tell me about yourself.")
    g = intro["guide"]
    assert [s["step"] for s in g["steps"]] == ["Present", "Past", "Why you fit", "Future"] and g["length"] and g["avoid"]
    assert any("Kevision Systems" in f for f in g["cv_facts"])                                   # facts come from the CV
    deep = next(q for q in qs if q["category"] == "Project deep dive")
    assert deep["guide"]["cv_facts"] and all(q.get("guide") for q in qs)

    assert client.post(f"/api/jobs/{job_id}/answer-draft", json={"question": intro["question"]}).status_code == 400  # no key
    from app.logic import ai
    monkeypatch.setattr(ai, "draft_answer", lambda *a, **k: "I'm Arpita… I also have 9 years of Kubernetes experience.")
    r = client.post(f"/api/jobs/{job_id}/answer-draft", json={"question": intro["question"]}).json()
    assert r["draft"].startswith("I'm Arpita") and any("Kubernetes" in w for w in r["warnings"])  # invented claim flagged
    assert client.post(f"/api/jobs/{job_id}/answer-draft", json={"question": "Unknown?"}).status_code == 404


def test_config_for_vercel_and_neon(monkeypatch):
    from app.config import DEFAULT_PASSWORD, DEFAULT_SECRET, Settings

    neon = Settings(DATABASE_URL="postgresql://u:p@ep-x-pooler.eu-central-1.aws.neon.tech/db?sslmode=require&channel_binding=require")
    assert neon.database_url.startswith("postgresql+psycopg://u:p@ep-x-pooler") and "sslmode=require" in neon.database_url
    monkeypatch.setenv("VERCEL", "1")
    unsafe = Settings(run_mode="auto", secret_key=DEFAULT_SECRET, master_password=DEFAULT_PASSWORD)
    assert unsafe.serverless
    with pytest.raises(RuntimeError, match="SECRET_KEY and MASTER_PASSWORD"):
        unsafe.check_for_production()
    Settings(run_mode="auto", secret_key="x" * 40, master_password="a-long-password").check_for_production()  # ok


def test_profile_from_json(monkeypatch):
    from app import defaults

    monkeypatch.setattr(defaults.settings, "profile_json", '{"name": "Jane Doe", "skills": ["Python", "SQL"], "target_roles": ["Data Engineer"]}')
    p = defaults.starting_profile()
    assert p["name"] == "Jane Doe" and p["skills"] == [{"name": "Python", "years": None}, {"name": "SQL", "years": None}]
    assert p["education"] == DEFAULT_PROFILE["education"]                    # missing keys fall back
    monkeypatch.setattr(defaults.settings, "profile_json", '{"nmae": "typo"}')
    with pytest.raises(ValueError, match="unknown keys: nmae"):
        defaults.starting_profile()


def test_serverless_steps_and_cron(client, monkeypatch):
    from app.config import settings

    client.patch("/api/settings", json={"sources": ["demo"], "search_queries": ["*", "Python"]})
    first = client.post("/api/jobs/search-step", json={"cursor": 0, "seconds": 0.001}).json()   # tiny budget: 1 search
    assert first["cursor"] == 1 and first["total"] == 2 and not first["done"] and first["new"] == 6
    rest = client.post("/api/jobs/search-step", json={"cursor": first["cursor"]}).json()
    assert rest["done"] and rest["duplicates"] >= 1                                              # resumes, no double saves

    monkeypatch.setattr(settings, "run_mode", "serverless")
    assert client.get("/api/info").json()["mode"] == "serverless"
    run = client.post("/api/automation/start", json={}).json()
    assert run["status"] == "RUNNING"                                                            # no thread started
    run = client.post("/api/automation/step").json()
    assert run["status"] == "FINISHED" and run["config"]["cursor"] == 2

    assert client.get("/api/cron/daily").status_code == 503                                      # CRON_SECRET not set
    monkeypatch.setattr(settings, "cron_secret", "s3cret")
    assert client.get("/api/cron/daily", headers={"Authorization": "Bearer nope"}).status_code == 401
    ok = client.get("/api/cron/daily", headers={"Authorization": "Bearer s3cret"}).json()
    assert ok["stepped"] == []            # today's run already finished: no second run the same day


def test_closed_jobs_bookmarks_deadline_and_source_health(client, monkeypatch):
    from app.services import freshness
    client.patch("/api/settings", json={"sources": ["demo"], "search_queries": ["*"]})
    client.post("/api/jobs/search-step", json={"cursor": 0})
    health = {s["source"]: s for s in client.get("/api/source-health").json()["sources"]}
    assert health["demo"]["last_ok_at"] and health["demo"]["jobs"] == 6 and health["demo"]["last_error"] is None

    jobs = client.get("/api/jobs", params={"view": "all"}).json()["items"]
    good = [j for j in jobs if j["ats_score"] >= 70]
    gone = good[0]["source_url"].rsplit("/", 1)[1]
    # job 0: both of its links are gone → closed; every other job: its pages are fine
    monkeypatch.setattr(freshness, "is_closed", lambda url: url.split("/jobs/")[1].split("/")[0] == gone)
    monkeypatch.setattr(freshness.time, "sleep", lambda s: None)
    monkeypatch.setattr(freshness, "NOT_CHECKED", ("manual",))
    result = client.post("/api/jobs/recheck").json()
    assert result["checked"] == len(good) and [c["id"] for c in result["closed"]] == [good[0]["id"]]
    visible = {j["id"] for j in client.get("/api/jobs", params={"view": "all"}).json()["items"]}
    assert good[0]["id"] not in visible                                                   # closed → hidden (and never applied to)
    assert good[0]["id"] in {j["id"] for j in client.get("/api/jobs", params={"view": "all", "show_closed": True}).json()["items"]}
    monkeypatch.setattr(freshness, "is_closed", lambda url: not url.endswith("/apply"))         # one open link is enough
    client.patch(f"/api/jobs/{good[0]['id']}", json={"reopen": True})
    assert client.post("/api/jobs/recheck").json()["closed"] == []
    assert good[0]["id"] in {j["id"] for j in client.get("/api/jobs", params={"view": "all"}).json()["items"]}

    client.patch(f"/api/jobs/{good[1]['id']}", json={"is_saved": True, "deadline": "2026-11-30"})
    [saved] = client.get("/api/jobs", params={"view": "all", "saved": True}).json()["items"]
    assert saved["id"] == good[1]["id"] and saved["deadline"] == "2026-11-30"
    assert client.patch(f"/api/jobs/{good[1]['id']}", json={"deadline": "30.11.2026"}).status_code == 422

    # "not for me": hidden from the list and the dashboard, survives a re-scrape (which re-opens closed jobs), undo works
    target = good[1]["id"]
    was_matching = client.get(f"/api/jobs/{target}").json()["status"] == "QUALIFIED"
    before = client.get("/api/dashboard").json()["counts"]
    client.patch(f"/api/jobs/{target}", json={"dismissed": True})
    assert target not in {j["id"] for j in client.get("/api/jobs", params={"view": "all"}).json()["items"]}
    assert target not in {j["id"] for j in client.get("/api/dashboard").json()["top_matches"]}
    after = client.get("/api/dashboard").json()["counts"]
    key = next(k for k in after if k.startswith("Matching"))
    assert after[key] == before[key] - (1 if was_matching else 0)
    client.patch(f"/api/jobs/{target}", json={"close": True, "reopen": False})
    client.patch(f"/api/jobs/{target}", json={"reopen": True})  # like a scrape that sees it again
    assert target not in {j["id"] for j in client.get("/api/jobs", params={"view": "all"}).json()["items"]}
    client.patch(f"/api/jobs/{target}", json={"dismissed": False})
    assert target in {j["id"] for j in client.get("/api/jobs", params={"view": "all"}).json()["items"]}


def test_interviews_and_calendar(client):
    from datetime import UTC, datetime, timedelta
    job_id = client.post("/api/jobs/import", json=JOB).json()["id"]
    other = client.post("/api/jobs/import", json={**JOB, "title": "Python Developer", "company": "Other GmbH",
                                                   "description": "Python FastAPI developer. " + SCADA_JD[::-1],
                                                   "application_url": "https://example.com/jobs/other"}).json()["id"]
    app = client.post(f"/api/jobs/{job_id}/applied", json={}).json()
    soon = (datetime.now(UTC) + timedelta(days=3)).replace(microsecond=0)
    r = client.patch(f"/api/applications/{app['id']}", json={"interviews": [
        {"at": (soon + timedelta(days=7)).isoformat(), "kind": "Technical"}, {"at": soon.replace(tzinfo=None).isoformat(), "kind": "HR screen"}]})
    assert [i["kind"] for i in r.json()["interviews"]] == ["HR screen", "Technical"]          # sorted; bare time = UTC
    client.patch(f"/api/jobs/{other}", json={"deadline": (soon + timedelta(days=1)).date().isoformat()})
    events = client.get("/api/applications/calendar").json()
    assert [e["type"] for e in events] == ["follow_up", "interview", "deadline", "interview"]
    assert events[1]["title"].startswith("HR screen: SCADA") and events[2]["job_id"] == other


def test_cv_import_reads_word_and_pdf_text(client):
    import base64
    cv = client.get("/api/cv.docx").content
    r = client.post("/api/profile/cv-import", json={"filename": "my_cv.docx", "content_base64": base64.b64encode(cv).decode()}).json()
    assert "Kevision Systems" in r["text"] and r["draft"] is None and "AI_API_KEY" in r["message"]   # no AI key in tests
    bad = client.post("/api/profile/cv-import", json={"filename": "cv.txt", "content_base64": base64.b64encode(b"hello").decode()})
    assert bad.status_code == 422


def test_cv_import_with_claude_returns_a_draft_only(client, monkeypatch):
    import base64
    from app.logic import ai
    fake = {"headline": "H", "summary": "S", "experience": [], "projects": [], "languages": [],
            "education": [{"degree": "MSc Computer Science", "school": "Greenwich University", "dates": "January 2022 – January 2023"}],
            "skills": ["python", "OPC-UA"]}
    monkeypatch.setattr(ai, "available", lambda: True)
    monkeypatch.setattr(ai, "extract_cv", lambda text: dict(fake))
    before = client.get("/api/profile").json()
    r = client.post("/api/profile/cv-import", json={"filename": "cv.docx", "content_base64": base64.b64encode(client.get("/api/cv.docx").content).decode()}).json()
    assert r["draft"]["education"] == ["MSc Computer Science | Greenwich University | January 2022 – January 2023"]
    assert [s["name"] for s in r["draft"]["skills"]] == ["Python", "OPC UA"]
    assert client.get("/api/profile").json() == before                                     # nothing saved until you click Save


def test_recruitment_directory(client, monkeypatch):
    from app.data.recruitment_companies import RECRUITMENT_COMPANIES
    from app.services import agencies
    assert len(RECRUITMENT_COMPANIES) == 50
    assert len({c["name"] for c in RECRUITMENT_COMPANIES}) == 50 and len({c["seed_id"] for c in RECRUITMENT_COMPANIES}) == 50
    assert all(c["website_url"].startswith("https://") for c in RECRUITMENT_COMPANIES)
    data = client.get("/api/agencies").json()
    assert data["total"] == 50                                                   # all 50 seeded for the user
    first = data["items"][0]
    assert first["status"] == "NOT_REVIEWED" and first["verification_status"] == "unverified"
    assert all(a[f] is None for a in data["items"] for f in ("india_candidate_support_verified", "visa_sponsorship_support_verified",
                                                              "relocation_support_verified"))   # never assumed
    assert {a["name"] for a in client.get("/api/agencies", params={"category": "india"}).json()["items"]} >= {"IndiaWorks", "Velai"}
    assert [a["name"] for a in client.get("/api/agencies", params={"q": "hays"}).json()["items"]] == ["Hays"]
    assert all("SAP & ERP" in a["job_family_specializations"] for a in client.get("/api/agencies", params={"family": "SAP & ERP"}).json()["items"])
    assert client.get("/api/agencies/options").json()["statuses"][0] == "NOT_REVIEWED"

    hays = client.get("/api/agencies", params={"q": "hays"}).json()["items"][0]
    assert client.patch(f"/api/agencies/{hays['id']}", json={"status": "MAYBE"}).status_code == 422
    assert client.patch(f"/api/agencies/{hays['id']}", json={"visa_sponsorship_support_verified": True}).status_code == 422  # no source
    r = client.patch(f"/api/agencies/{hays['id']}", json={"status": "PROFILE_REGISTERED", "is_saved": True, "notes": "Registered online",
                                                         "relocation_support_verified": True, "source_url": "https://www.hays.de/en/relocation"}).json()
    assert r["verification_status"] == "confirmed" and r["relocation_support_verified"] is True and r["visa_sponsorship_support_verified"] is None
    assert [a["name"] for a in client.get("/api/agencies", params={"saved": True}).json()["items"]] == ["Hays"]

    # a job advertised by ALTEN is linked to the recruiter; applying saves the application AND sets the recruiter to Applied
    job = client.post("/api/jobs/import", json={**JOB, "company": "ALTEN GmbH", "source": "smartrecruiters"}).json()
    assert client.get(f"/api/jobs/{job['id']}").json()["agency"] == "ALTEN Germany"
    alten = client.get("/api/agencies", params={"q": "alten"}).json()["items"][0]
    assert alten["jobs"] == 1 and alten["ingestion"] == "automatic"
    assert client.get(f"/api/agencies/{alten['id']}/jobs").json()[0]["employer"] is None   # ad names no other employer
    client.post(f"/api/jobs/{job['id']}/applied", json={})
    alten = client.get("/api/agencies", params={"q": "alten"}).json()["items"][0]
    assert alten["status"] == "APPLIED" and alten["applied"] == 1 and alten["last_contact_at"]
    assert [a["job"]["company"] for a in client.get("/api/applications").json()] == ["ALTEN GmbH"]

    monkeypatch.setattr(agencies, "check_website", lambda url: ("website_ok", "HTTP 200") if "hays" not in url else ("website_blocked", "HTTP 403"))
    monkeypatch.setattr(agencies.time, "sleep", lambda s: None)
    result = client.post("/api/agencies/check-websites").json()
    assert result["checked"] == 50 and result["ok"] == 49
    hays = client.get("/api/agencies", params={"q": "hays"}).json()["items"][0]
    assert hays["verification_status"] == "confirmed"                            # your confirmation is never downgraded
    assert client.get("/api/agencies", params={"q": "velai"}).json()["items"][0]["verification_status"] == "website_ok"


def test_agency_matching_is_exact():
    from app.services.agencies import _key
    assert _key("ALTEN GmbH") == _key("Alten") == "alten"
    assert _key("Hays AG") == "hays" and _key("Haystack Software") != "hays"
